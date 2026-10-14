# -*- coding: utf-8 -*-
"""V23 最终十维度验证（修正版）。"""
import io, re, json, zipfile, collections
from docx import Document

P = r"E:\deepseekworkspace\manuscript_JCP_v23.docx"
d = Document(P)
ps = [p.text.strip() for p in d.paragraphs]
full = "\n".join(ps)
R = {}


def chk(tag, ok, detail=""):
    R[tag] = bool(ok)
    print("  %-42s %s  %s" % (tag, "PASS" if ok else "FAIL", detail))


print("=" * 100)
print("V23 十维度严格检查（最终）")
print("=" * 100)

figs = [int(re.match(r"^Figure (\d+) \|", t).group(1)) for t in ps if re.match(r"^Figure \d+ \|", t)]
tabs = [int(re.match(r"^Table (\d+) \|", t).group(1)) for t in ps if re.match(r"^Table \d+ \|", t)]

# D1
REQ = ["Abstract", "1. Introduction", "2. Results", "3. Conclusions", "4. Methods",
       "Data availability", "Code availability", "Competing interests", "Funding", "References"]
pos = [next((i for i, t in enumerate(ps) if t == k), None) for k in REQ]
chk("D1 结构：必需章节齐全且顺序正确",
    all(p is not None for p in pos) and pos == sorted(x for x in pos if x is not None))

# D2
chk("D2 编号：图 1..10 连续", figs == list(range(1, 11)), str(figs))
chk("D2 编号：表 1..6 连续", tabs == list(range(1, 7)), str(tabs))
eqs = re.findall(r"\((\d+)\)\s*$", full, re.M)
chk("D2 编号：公式 1..5 存在", all(("(%d)" % k) in full for k in range(1, 6)))

# D3
caps = {int(re.match(r"^Figure (\d+) \|", t).group(1)): i for i, t in enumerate(ps)
        if re.match(r"^Figure \d+ \|", t)}
unc_f = [n for n, i in caps.items()
         if not any(re.search(r"\b(Fig\.|Figure)\s*%d\b" % n, ps[j]) for j in range(len(ps)) if j != i)]
chk("D3 交叉引用：每张图都被引用", not unc_f, "未引用=%s" % (unc_f or "无"))
capt = {int(re.match(r"^Table (\d+) \|", t).group(1)): i for i, t in enumerate(ps)
        if re.match(r"^Table \d+ \|", t)}
unc_t = [n for n, i in capt.items()
         if not any(re.search(r"\bTable\s*%d\b" % n, ps[j]) for j in range(len(ps)) if j != i)]
chk("D3 交叉引用：每张表都被引用", not unc_t, "未引用=%s" % (unc_t or "无"))
dang = [m.group(0) for m in re.finditer(r"\b(?:Fig\.|Figure)\s*(\d+)", full)
        if int(m.group(1)) not in caps]
chk("D3 交叉引用：无悬空引用", not dang, "悬空=%s" % (sorted(set(dang)) or "无"))

# D4
iRef = ps.index("References")
refs = [int(m.group(1)) for t in ps[iRef + 1:] for m in [re.match(r"^(\d+)\.\s", t)] if m]
cited = set()
for m in re.finditer(r"REF:(\d+(?:,\d+)*)", full):
    for x in m.group(1).split(","):
        cited.add(int(x))
chk("D4 参考文献 1..33 连续", refs == list(range(1, 34)), "n=%d" % len(refs))
chk("D4 无越界引用", not [x for x in cited if x < 1 or x > 33], "引用数=%d" % len(cited))

# D5
bad = [(ti + 1, ri) for ti, tb in enumerate(d.tables) if len(tb.rows) > 1
       for ri, r in enumerate(tb.rows) if any(not c.text.strip() for c in r.cells)]
chk("D5 表格：无空单元格", not bad, "%d 处" % len(bad))

# D6
ai = ps.index("Abstract")
abs_txt = ps[ai + 1]
wa = len(re.findall(r"\S+", abs_txt))
chk("D6 摘要 <=250 词", wa <= 250, "%d 词" % wa)
chk("D6 摘要无引用", not re.search(r"\[|REF:", abs_txt))
chk("D6 关键词存在", any(t.startswith("Keywords") for t in ps))

# D7
for pat, tag in [(r"@@", "占位符"), (r"Supplementary", "Supplementary"), (r"REF:\d", "REF 标记"),
                 (r"XXX|TODO|TBD", "模板标记")]:
    h = [t[:60] for t in ps if re.search(pat, t)]
    chk("D7 文本卫生：无%s残留" % tag, not h, str(h[:2]) if h else "")

# D8
z = zipfile.ZipFile(P)
media = [n for n in z.namelist() if "media/" in n]
chk("D8 图形资源：嵌入图片 >= 图注数", len(media) >= len(figs), "%d 图 / %d 图注" % (len(media), len(figs)))

# D9
KEY = ["0.6 percent", "4 to 82 percent", "16 of 18", "4 of 14", "48 to 108", "2e-14"]
chk("D9 数字溯源：关键数字在文中", all(s in full for s in KEY),
    "缺=%s" % [s for s in KEY if s not in full])

# D10
TOP = {1: "decay classes", 2: "rank", 3: "Recovering", 4: "rank laws", 5: "Diagnostics",
       6: "Controlled operators", 7: "two exponents", 8: "alignment", 9: "Estimator",
       10: "spectral branch"}
mis = [n for n, k in TOP.items()
       if not any(k.lower() in t.lower() for t in ps if re.match(r"^Figure %d \|" % n, t))]
chk("D10 一致性：图注主题与序号匹配", not mis, "不匹配=%s" % (mis or "无"))

print("=" * 100)
np = sum(1 for v in R.values() if v)
print("汇总：%d/%d PASS" % (np, len(R)))
for k, v in R.items():
    if not v:
        print("   FAIL: %s" % k)
json.dump(R, io.open(r"E:\deepseekworkspace\v23_audit_final.json", "w", encoding="utf-8"),
          indent=2, ensure_ascii=False)
