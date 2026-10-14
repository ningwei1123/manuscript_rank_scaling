# -*- coding: utf-8 -*-
"""扫描全文中真正需要语言改进的段落（冗余/长句/被动/重复）。"""
import io, re
from docx import Document

P = r"E:\deepseekworkspace\manuscript_JCP_v49.docx"
d = Document(P)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def fe(el):
    return "".join((n.text or "") for n in el.iter() if n.tag in (W + "t", M + "t"))


print("=" * 108)
print("① 超长句（>55 词）")
print("=" * 108)
for i, p in enumerate(d.paragraphs):
    t = fe(p._element).strip()
    if not t or re.match(r'^(Fig\.|Table) \d+ \|', t):
        continue
    for s in re.split(r'(?<=[.])\s+', t):
        w = len(s.split())
        if w > 55:
            print(f"\n[段 {i}] {w} 词: {s[:230]}")

print("\n" + "=" * 108)
print("② 冗余/松散表达")
print("=" * 108)
PAT = [
 (r'\bin order to\b', 'in order to'),
 (r'\bdue to the fact that\b', 'due to the fact that'),
 (r'\bit is worth (saying|noting|stating)\b', 'it is worth saying'),
 (r'\bwe (can|may) note that\b', 'we can note that'),
 (r'\bIt is (also )?(important|worth)\b', 'It is important'),
 (r'\bis not the same as\b', 'is not the same as'),
 (r'\bthe fact that\b', 'the fact that'),
 (r'\bthere (is|are) (a|an|no)\b', 'there is a'),
 (r'\bwhich is why\b', 'which is why'),
 (r'\bwe do not claim\b', 'we do not claim'),
 (r'\bof course\b', 'of course'),
 (r'\bvery (nearly|close)\b', 'very nearly'),
 (r'\bThis is (not )?because\b', 'This is because'),
 (r'\bat all\b', 'at all'),
 (r'\balso\b.*\balso\b', 'also...also'),
]
for pat, lab in PAT:
    hits = []
    for i, p in enumerate(d.paragraphs):
        t = fe(p._element).strip()
        for m in re.finditer(pat, t, re.I):
            hits.append((i, t[max(0, m.start() - 60):m.end() + 60]))
    if hits:
        print(f"\n[{lab}] {len(hits)} 处")
        for i, ctx in hits[:3]:
            print(f"   [段 {i}] ...{ctx}...")

print("\n" + "=" * 108)
print("③ 被动语态密集句")
print("=" * 108)
for i, p in enumerate(d.paragraphs):
    t = fe(p._element).strip()
    if len(t.split()) < 60:
        continue
    npas = len(re.findall(r'\b(is|are|was|were|be|been|being)\s+\w+(ed|en)\b', t))
    if npas >= 5:
        print(f"  [段 {i}] 被动 {npas} 处，{len(t.split())} 词: {t[:130]}")
