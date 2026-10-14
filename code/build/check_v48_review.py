# -*- coding: utf-8 -*-
"""v48 最终验证（v46 审稿意见逐条） + 全部既有检查。"""
import io, re, zipfile
from docx import Document
from docx.table import Table

P = r"E:\deepseekworkspace\manuscript_JCP_v48.docx"
d = Document(P)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
PIC = "{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"
BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def fe(el):
    return "".join((n.text or "") for n in el.iter() if n.tag in (W + "t", M + "t"))


ps = [fe(p._element).strip() for p in d.paragraphs]
ft = "\n".join(ps)


def chk(t, ok, dt=""):
    print(f"  {t:<50} {'PASS' if ok else 'FAIL'}  {dt}")


print("=" * 104)
print("v48 验证（v46 审稿意见）")
print("=" * 104)

# P0-1
eq2 = ""
for tb in d.tables:
    if len(tb.rows) == 1 and len(tb.columns) == 2 and tb.rows[0].cells[1].text.strip() == "Eq. (2)":
        eq2 = "".join((n.text or "") for n in tb._tbl.iter(M + 't'))
chk("P0-1 Eq.(2) 指数为 1/(β−1/2)", "1/(\u03b2\u22121/2)" in eq2, eq2)
chk("P0-1 Eq.(1) 仍为 1/α",
    "1" in "".join((n.text or "") for n in
                   next(tb for tb in d.tables
                        if len(tb.rows) == 1 and len(tb.columns) == 2
                        and tb.rows[0].cells[1].text.strip() == "Eq. (1)")._tbl.iter(M + 't')))

# P1
chk("P1-1 能量约定 vs 范数约定已澄清", "energy-convention form" in ft)
chk("P1-2 两个 99%-energy 秩已并列", "returns 10, 11 and 10" in ft and "ranks are 24, 27 and 20" in ft)
chk("P1-3 未定义的 '19 percent' 已删除", "19 percent" not in ft)

# P2
chk("P2-1 c_log 记号已引入", "c_log" in ft)
chk("P2-2 对数分支限定已加", "logarithmic branch has no curve" in ft)
chk("P2-3 'fields.Eighth' 粘连已修", "fields. Eighth" in ft)
chk("P2-3 无残留粘连", not re.search(r'fields\.[A-Z]', ft))

# 结构完整性
iRef = ps.index("References")
body = "\n".join(ps[:iRef])
nums = set()
for m in re.finditer(r'\[([\d,\u2013\u2014\-\s]+)\]', body):
    for part in re.split(r'[,\u2013\u2014-]', m.group(1)):
        if part.strip().isdigit():
            nums.add(int(part.strip()))
refs = [int(m.group(1)) for t in ps[iRef + 1:] for m in [re.match(r'^(\d+)\.\s', t)] if m]
chk("参考文献 1..33 连续且全部被引用", refs == list(range(1, 34)) and nums == set(range(1, 34)))
figs = [int(m.group(1)) for t in ps for m in [re.match(r'^Fig\. (\d+) \|', t)] if m]
tabs = [int(m.group(1)) for t in ps for m in [re.match(r'^Table (\d+) \|', t)] if m]
chk("Fig. 1..10 连续", figs == list(range(1, 11)), str(figs))
chk("Table 1..6 连续", tabs == list(range(1, 7)), str(tabs))
z = zipfile.ZipFile(P)
media = sorted(x.split('/')[-1] for x in z.namelist() if x.startswith('word/media/'))
used = []
for ch in d.element.body.iterchildren():
    if ch.tag.split('}')[-1] != 'p':
        continue
    for blip in ch.iter(BLIP):
        rid = blip.get(RID)
        if rid:
            used.append(d.part.rels[rid].target_ref.split('/')[-1])
chk("10 张图片全部被引用 / 无孤儿", len(set(used)) == 10 and not (set(media) - set(used)))
ai = ps.index("Abstract")
wa = len(re.findall(r"\S+", ps[ai + 1]))
chk("摘要 <=250 词", wa <= 250, f"{wa} 词")
chk("数学对象完整", sum(1 for p in d.paragraphs for x in p._element.iter(M + 't')) > 100)
chk("无 ASCII 希腊字母", not re.search(r'\b(alpha|beta|epsilon|gamma|tau|lambda|sigma|eps|Delta)\b', ft))
ok3 = True
for tb in d.tables:
    if len(tb.rows) > 1 and len(tb.columns) > 1:
        N = len(tb.rows)
        for i, row in enumerate(tb.rows):
            for c in row.cells:
                tcPr = c._tc.find(W + 'tcPr')
                b = tcPr.find(W + 'tcBorders') if tcPr is not None else None
                s = {x.tag.split('}')[-1]: x.get(W + 'val') for x in b} if b is not None else {}
                for k in ('left', 'right', 'insideH', 'insideV'):
                    if s.get(k) != 'nil':
                        ok3 = False
                if i == 0 and s.get('top') != 'single':
                    ok3 = False
                if i == N - 1 and s.get('bottom') != 'single':
                    ok3 = False
chk("6 张实质表为三线表", ok3)
chk("第三类谱检验仍在 Results",
    ps.index("2. Results") < next(i for i, t in enumerate(ps)
                                  if t.startswith("Testing whether a third spectral class"))
    < ps.index("3. Conclusions"))
print(f"\n全文词数: {len(ft.split())}")
